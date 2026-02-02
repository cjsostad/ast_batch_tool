"""
Failed Job Tracker Module
Tracks jobs that fail during batch processing and creates a spreadsheet for later retry.
"""

import os
import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


class FailedJobTracker:
    """Tracks and logs failed jobs to a separate spreadsheet for later retry"""
    
    def __init__(self, output_directory, logger=None):
        """
        Initialize the failed job tracker
        
        Args:
            output_directory: Directory where the failed jobs spreadsheet will be saved
            logger: Logger instance for logging
        """
        self.output_directory = output_directory
        self.logger = logger
        self.failed_jobs_file = None
        self.failed_jobs = []
        
    def initialize_failed_jobs_spreadsheet(self):
        """Create the failed jobs spreadsheet with headers"""
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        self.failed_jobs_file = os.path.join(
            self.output_directory, 
            f"failed_jobs_{timestamp}.xlsx"
        )
        
        # Create new workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Failed Jobs"
        
        # Define headers
        headers = [
            'spreadsheet_name',
            'job_index',
            'failure_type',
            'failure_time',
            'region',
            'feature_layer',
            'crown_file_number',
            'disposition_number',
            'parcel_number',
            'output_directory',
            'notes'
        ]
        
        # Write headers with formatting
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 20  # spreadsheet_name
        ws.column_dimensions['B'].width = 12  # job_index
        ws.column_dimensions['C'].width = 15  # failure_type
        ws.column_dimensions['D'].width = 20  # failure_time
        ws.column_dimensions['E'].width = 15  # region
        ws.column_dimensions['F'].width = 40  # feature_layer
        ws.column_dimensions['G'].width = 20  # crown_file_number
        ws.column_dimensions['H'].width = 20  # disposition_number
        ws.column_dimensions['I'].width = 15  # parcel_number
        ws.column_dimensions['J'].width = 40  # output_directory
        ws.column_dimensions['K'].width = 50  # notes
        
        # Save the workbook
        wb.save(self.failed_jobs_file)
        
        if self.logger:
            self.logger.info(f"Failed Job Tracker: Created spreadsheet at {self.failed_jobs_file}")
        print(f"Failed Job Tracker: Created spreadsheet at {self.failed_jobs_file}")
        
        return self.failed_jobs_file
    
    def log_failed_job(self, spreadsheet_name, job_index, job_data, failure_type, notes=""):
        """
        Log a failed job to the tracking spreadsheet
        
        Args:
            spreadsheet_name: Name of the Excel file being processed (e.g., 'jobs_15.xlsx')
            job_index: Index of the job that failed
            job_data: Dictionary containing job parameters
            failure_type: Type of failure ('timeout', 'crash', 'worker_error', 'unknown')
            notes: Additional notes about the failure
        """
        # Initialize spreadsheet if not already done
        if self.failed_jobs_file is None:
            self.initialize_failed_jobs_spreadsheet()
        
        # Load the workbook
        wb = load_workbook(self.failed_jobs_file)
        ws = wb.active
        
        # Prepare row data
        failure_time = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        row_data = [
            spreadsheet_name,
            job_index,
            failure_type,
            failure_time,
            job_data.get('region', ''),
            job_data.get('feature_layer', ''),
            job_data.get('crown_file_number', ''),
            job_data.get('disposition_number', ''),
            job_data.get('parcel_number', ''),
            job_data.get('output_directory', ''),
            notes
        ]
        
        # Append to the spreadsheet
        ws.append(row_data)
        
        # Apply color coding based on failure type
        row_num = ws.max_row
        if failure_type == 'timeout':
            fill_color = "FFCCCC"  # Light red
        elif failure_type == 'crash':
            fill_color = "FFDDAA"  # Light orange
        elif failure_type == 'worker_error':
            fill_color = "FFFFCC"  # Light yellow
        else:
            fill_color = "DDDDDD"  # Light gray
        
        for col_idx in range(1, len(row_data) + 1):
            ws.cell(row=row_num, column=col_idx).fill = PatternFill(
                start_color=fill_color, end_color=fill_color, fill_type="solid"
            )
        
        # Save the workbook
        wb.save(self.failed_jobs_file)
        
        # Log the failure
        if self.logger:
            self.logger.error(
                f"Failed Job Tracker: Logged failure - {spreadsheet_name} job {job_index} "
                f"({failure_type}): {notes}"
            )
        print(
            f"Failed Job Tracker: Logged failure - {spreadsheet_name} job {job_index} "
            f"({failure_type})"
        )
        
        # Add to in-memory list
        self.failed_jobs.append({
            'spreadsheet': spreadsheet_name,
            'job_index': job_index,
            'failure_type': failure_type,
            'job_data': job_data
        })
    
    def get_failed_jobs_count(self):
        """Return the total number of failed jobs logged"""
        return len(self.failed_jobs)
    
    def get_failed_jobs_file(self):
        """Return the path to the failed jobs spreadsheet"""
        return self.failed_jobs_file
    
    def print_summary(self):
        """Print a summary of failed jobs"""
        if not self.failed_jobs:
            print("Failed Job Tracker: No failed jobs to report")
            if self.logger:
                self.logger.info("Failed Job Tracker: No failed jobs to report")
            return
        
        timeout_count = sum(1 for j in self.failed_jobs if j['failure_type'] == 'timeout')
        crash_count = sum(1 for j in self.failed_jobs if j['failure_type'] == 'crash')
        worker_error_count = sum(1 for j in self.failed_jobs if j['failure_type'] == 'worker_error')
        unknown_count = sum(1 for j in self.failed_jobs if j['failure_type'] == 'unknown')
        
        summary = f"""
Failed Job Tracker Summary:
===========================
Total Failed Jobs: {len(self.failed_jobs)}
  - Timeouts: {timeout_count}
  - Crashes: {crash_count}
  - Worker Errors: {worker_error_count}
  - Unknown: {unknown_count}

Failed jobs saved to: {self.failed_jobs_file}
"""
        print(summary)
        if self.logger:
            self.logger.info(summary)
